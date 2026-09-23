import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from tools import BuildWord


class BuildWordMeetingReportTests(unittest.TestCase):
    def test_tool_formats_meeting_report_with_decisions_and_action_table(self):
        body = """RESUME_COURT:
Le calendrier de lancement est validé.
COMPTE_RENDU_DETAILLE:
1. Résumé long
Le projet a été présenté à l'équipe.
2. Points importants
- Le lancement est prévu en octobre.
3. Décisions prises
- Le calendrier est approuvé.
4. Actions à faire
- Action : Envoyer le planning | Responsable : Awa | Échéance : 15 octobre
5. Questions ouvertes
Non précisé
6. Risques ou blocages
Non précisé"""

        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"REPORTS_DIR": directory}):
            result = BuildWord().forward(
                title="Réunion de lancement",
                date="20/09/2026 09:00",
                body=body,
                filename="meeting-42-lancement",
                template="meeting_report",
                organization="Génuka",
                participants="Awa, Boris",
            )

            path = Path(directory) / "meeting-42-lancement.docx"
            self.assertTrue(path.exists(), result)
            document = Document(path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("COMPTE RENDU DE RÉUNION", text)
            self.assertIn("Génuka", text)
            self.assertIn("Le calendrier de lancement est validé.", text)
            self.assertIn("Le calendrier est approuvé.", text)
            self.assertEqual(["Action", "Responsable", "Échéance"], [cell.text for cell in document.tables[0].rows[0].cells])
            self.assertEqual(["Envoyer le planning", "Awa", "15 octobre"], [cell.text for cell in document.tables[0].rows[1].cells])
            self.assertIn("RAI-00042", document.sections[0].footer.paragraphs[0].text)
            self.assertNotIn("Transcription source", text)

    def test_meeting_filename_uses_template_even_if_agent_omits_option(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"REPORTS_DIR": directory}):
            result = BuildWord().forward(
                title="Suivi",
                body="Les participants ont discuté du calendrier.",
                filename="meeting-8-suivi",
            )
            document = Document(Path(directory) / "meeting-8-suivi.docx")
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("COMPTE RENDU DE RÉUNION", text, result)
            self.assertIn("Les participants ont discuté du calendrier.", text)
            self.assertEqual([], document.tables)

    def test_other_documents_keep_generic_word_format(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"REPORTS_DIR": directory}):
            result = BuildWord().forward(title="Lettre de suivi", body="Bonjour.", filename="lettre")
            document = Document(Path(directory) / "lettre.docx")
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Lettre de suivi", text, result)
            self.assertNotIn("COMPTE RENDU DE RÉUNION", text)


if __name__ == "__main__":
    unittest.main()
