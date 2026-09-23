from __future__ import annotations
import os
import json
import re
import unicodedata
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape
import json
from smolagents import Tool


def get_reports_dir() -> Path:
    reports_dir = Path(os.getenv("REPORTS_DIR", "/app/storage/reports"))
    reports_dir.mkdir(parents=True, exist_ok=True)
    return reports_dir


class BuildWord(Tool):
    name = "BuildWord"
    description = (
        "Crée un document Word professionnel (.docx), dont des comptes rendus de réunion "
        "avec le modèle meeting_report. "
        "Peut générer des lettres, CV, rapports avec tableaux simples ou complexes. "
        "Supporte la fusion de cellules, styles personnalisés, et mise en page avancée."
    )

    inputs = {
        "title": {"type": "string", "description": "Titre principal du document"},
        "recipient": {
            "type": "string",
            "description": "Destinataire (optionnel, multiligne)",
            "nullable": True,
        },
        "sender": {
            "type": "string",
            "description": "Expéditeur (optionnel, multiligne)",
            "nullable": True,
        },
        "date": {"type": "string", "description": "Date (optionnel)", "nullable": True},
        "subject": {
            "type": "string",
            "description": "Objet (optionnel)",
            "nullable": True,
        },
        "body": {
            "type": "string",
            "description": "Corps du document. Supporte :\\n- Paragraphes (séparés par \\n\\n)\\n- Listes (lignes commençant par '- ')\\n- Tableaux (JSON: [TABLE]json_data[/TABLE])",
            "nullable": True,
        },
        "tables": {
            "type": "string",
            "description": "JSON array de tableaux complexes (optionnel). Format: [{headers:[...], rows:[[...]], merge_cells:[...], styles:{...}}]",
            "nullable": True,
        },
        "filename": {
            "type": "string",
            "description": "Nom du fichier sans extension,",
            "nullable": True,
        },
        "style_config": {
            "type": "string",
            "description": "JSON de config style (optionnel): {font:'Arial', font_size:11, margins:2.5, line_spacing:1.15}",
            "nullable": True,
        },
        "template": {
            "type": "string",
            "description": "Utilise meeting_report pour un compte rendu de réunion structuré (optionnel)",
            "nullable": True,
        },
        "organization": {
            "type": "string",
            "description": "Nom de l'organisation pour le compte rendu (optionnel)",
            "nullable": True,
        },
        "participants": {
            "type": "string",
            "description": "Participants à la réunion (optionnel)",
            "nullable": True,
        },
    }

    output_type = "string"

    def _set_cell_border(self, cell, **kwargs):
        """Ajoute des bordures personnalisées à une cellule"""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")

        for edge in ("top", "left", "bottom", "right"):
            if edge in kwargs:
                edge_data = kwargs[edge]
                edge_el = OxmlElement(f"w:{edge}")
                edge_el.set(qn("w:val"), edge_data.get("val", "single"))
                edge_el.set(qn("w:sz"), str(edge_data.get("sz", 4)))
                edge_el.set(qn("w:space"), str(edge_data.get("space", 0)))
                edge_el.set(qn("w:color"), edge_data.get("color", "000000"))
                tcBorders.append(edge_el)

        tcPr.append(tcBorders)

    def _apply_cell_style(self, cell, style_config):
        """Applique un style à une cellule"""
        if "background" in style_config:
            shading_elm = OxmlElement("w:shd")
            shading_elm.set(qn("w:fill"), style_config["background"])
            cell._tc.get_or_add_tcPr().append(shading_elm)

        if (
            "bold" in style_config
            or "font_size" in style_config
            or "color" in style_config
        ):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if style_config.get("bold"):
                        run.font.bold = True
                    if "font_size" in style_config:
                        run.font.size = Pt(style_config["font_size"])
                    if "color" in style_config:
                        color = style_config["color"]
                        run.font.color.rgb = RGBColor(
                            int(color[0:2], 16),
                            int(color[2:4], 16),
                            int(color[4:6], 16),
                        )

        if "alignment" in style_config:
            for paragraph in cell.paragraphs:
                alignment_map = {
                    "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    "left": WD_ALIGN_PARAGRAPH.LEFT,
                    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                }
                paragraph.alignment = alignment_map.get(
                    style_config["alignment"], WD_ALIGN_PARAGRAPH.LEFT
                )

        if "vertical_alignment" in style_config:
            valign_map = {
                "center": WD_ALIGN_VERTICAL.CENTER,
                "top": WD_ALIGN_VERTICAL.TOP,
                "bottom": WD_ALIGN_VERTICAL.BOTTOM,
            }
            cell.vertical_alignment = valign_map.get(
                style_config["vertical_alignment"], WD_ALIGN_VERTICAL.TOP
            )

    def _create_table(self, doc, table_config):
        """Crée un tableau avec configuration avancée"""
        headers = table_config.get("headers", [])
        rows = table_config.get("rows", [])
        merge_cells = table_config.get("merge_cells", [])
        styles = table_config.get("styles", {})

        # Créer le tableau
        num_cols = len(headers) if headers else (len(rows[0]) if rows else 1)
        num_rows = len(rows) + (1 if headers else 0)
        table = doc.add_table(rows=num_rows, cols=num_cols)

        # Style général du tableau
        table.style = styles.get("table_style", "Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # En-têtes
        if headers:
            header_cells = table.rows[0].cells
            for i, header in enumerate(headers):
                header_cells[i].text = str(header)
                # Style header par défaut
                header_style = styles.get(
                    "header_style",
                    {
                        "bold": True,
                        "background": "D9E2F3",
                        "alignment": "center",
                        "vertical_alignment": "center",
                    },
                )
                self._apply_cell_style(header_cells[i], header_style)

        # Données
        start_row = 1 if headers else 0
        for i, row_data in enumerate(rows):
            row_cells = table.rows[start_row + i].cells
            for j, cell_data in enumerate(row_data):
                row_cells[j].text = str(cell_data)

                # Style des cellules de données
                if "cell_styles" in styles and f"{i},{j}" in styles["cell_styles"]:
                    self._apply_cell_style(
                        row_cells[j], styles["cell_styles"][f"{i},{j}"]
                    )
                elif "data_style" in styles:
                    self._apply_cell_style(row_cells[j], styles["data_style"])

        # Fusion de cellules
        for merge in merge_cells:
            start_row = merge.get("start_row", 0)
            start_col = merge.get("start_col", 0)
            end_row = merge.get("end_row", start_row)
            end_col = merge.get("end_col", start_col)

            if start_row != end_row or start_col != end_col:
                start_cell = table.rows[start_row].cells[start_col]
                end_cell = table.rows[end_row].cells[end_col]
                start_cell.merge(end_cell)

        # Largeur des colonnes
        if "column_widths" in styles:
            for i, width in enumerate(styles["column_widths"]):
                for row in table.rows:
                    row.cells[i].width = Inches(width)

        return table

    def _parse_body_with_tables(self, doc, body_text):
        """Parse le corps du texte et insère les tableaux inline"""
        import re

        # Chercher les tableaux dans le texte
        table_pattern = r"\[TABLE\](.*?)\[/TABLE\]"
        parts = re.split(table_pattern, body_text, flags=re.DOTALL)

        for i, part in enumerate(parts):
            if i % 2 == 0:  # Texte normal
                self._add_text_content(doc, part)
            else:  # Tableau JSON
                try:
                    table_config = json.loads(part.strip())
                    self._create_table(doc, table_config)
                except json.JSONDecodeError as e:
                    doc.add_paragraph(f"[Erreur tableau: {str(e)}]")

    def _add_text_content(self, doc, text):
        """Ajoute du contenu texte (paragraphes et listes)"""
        if not text.strip():
            return

        paragraphs = text.split("\n\n")
        for para_text in paragraphs:
            para_text = para_text.strip()
            if not para_text:
                continue

            # Liste à puces
            if "- " in para_text and para_text.split("\n")[0].strip().startswith("- "):
                lines = para_text.split("\n")
                for line in lines:
                    line = line.strip()
                    if line.startswith("- "):
                        doc.add_paragraph(line[2:].strip(), style="List Bullet")
                    elif line:
                        doc.add_paragraph(line)
            else:
                doc.add_paragraph(para_text)

    @staticmethod
    def _meeting_heading_name(line):
        heading = line.strip().strip("#* ").rstrip(":*").strip()
        heading = re.sub(r"^\d+[.)]\s*", "", heading)
        heading = unicodedata.normalize("NFKD", heading).encode("ascii", "ignore").decode("ascii")
        heading = re.sub(r"\s+", " ", heading.lower())
        return {
            "resume long": "discussion",
            "resume detaille": "discussion",
            "points importants": "points",
            "points abordes": "points",
            "decisions prises": "decisions",
            "decisions": "decisions",
            "actions a faire": "actions",
            "plan d actions": "actions",
            "questions ouvertes": "questions",
            "points en suspens": "questions",
            "risques ou blocages": "risks",
            "risques et blocages": "risks",
        }.get(heading)

    def _meeting_content(self, body):
        """Sépare les rubriques textuelles fournies à BuildWord par l'agent."""
        short = re.search(
            r"RESUME_COURT\s*:\s*(.*?)(?=COMPTE_RENDU_DETAILLE\s*:|$)",
            body,
            flags=re.IGNORECASE | re.DOTALL,
        )
        detailed = re.search(
            r"COMPTE_RENDU_DETAILLE\s*:\s*(.*?)(?=WORD_PATH\s*:|$)",
            body,
            flags=re.IGNORECASE | re.DOTALL,
        )
        summary = short.group(1).strip() if short else ""
        report = detailed.group(1).strip() if detailed else "" if short else body
        sections = {}
        introduction = []
        current = None
        for line in report.splitlines():
            heading = self._meeting_heading_name(line)
            if heading:
                current = heading
                sections.setdefault(heading, [])
            elif current:
                sections[current].append(line)
            else:
                introduction.append(line)
        if introduction:
            sections["discussion"] = introduction + sections.get("discussion", [])
        return summary, {name: "\n".join(lines).strip() for name, lines in sections.items()}

    @staticmethod
    def _meeting_missing(text):
        plain = unicodedata.normalize("NFKD", (text or "").strip(" .*-"))
        plain = plain.encode("ascii", "ignore").decode("ascii").lower()
        return plain in {"", "non precise", "aucun", "aucune"}

    def _add_meeting_lines(self, container, content):
        if self._meeting_missing(content):
            container.add_paragraph("Non précisé dans la transcription.")
            return
        for raw_line in content.splitlines():
            line = raw_line.strip().strip("* ")
            if not line:
                continue
            bullet = re.match(r"^[-•]\s+(.+)$", line)
            numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
            if bullet:
                container.add_paragraph(bullet.group(1).strip(), style="List Bullet")
            elif numbered:
                container.add_paragraph(numbered.group(1).strip(), style="List Number")
            else:
                container.add_paragraph(line)

    @staticmethod
    def _meeting_action_rows(content):
        rows = []
        for raw_line in content.splitlines():
            line = re.sub(r"^\s*(?:[-•*]|\d+[.)])\s*", "", raw_line).strip()
            if not line:
                continue
            fields = {}
            for part in line.split("|"):
                part = part.strip().strip("*")
                match = re.match(
                    r"^(action|tâche|tache|responsable|échéance|echeance)\s*:\s*(.*)$",
                    part,
                    flags=re.IGNORECASE,
                )
                if match:
                    name = unicodedata.normalize("NFKD", match.group(1)).encode("ascii", "ignore").decode("ascii").lower()
                    key = "action" if name in {"action", "tache"} else "deadline" if name == "echeance" else "owner"
                    fields[key] = match.group(2).strip()
                elif part and "action" not in fields:
                    fields["action"] = part
            action = fields.get("action", "")
            if action and not BuildWord._meeting_missing(action):
                rows.append((action, fields.get("owner") or "Non précisé", fields.get("deadline") or "Non précisée"))
        return rows

    @staticmethod
    def _meeting_section(doc, title):
        paragraph = doc.add_paragraph(title, style="Heading 1")
        properties = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "6"), ("space", "7"), ("color", "D7E2F3")):
            bottom.set(qn(f"w:{key}"), value)
        borders.append(bottom)
        properties.append(borders)

    def _render_meeting_report(self, doc, title, date, body, filename, organization, participants):
        """Applique la présentation du compte rendu dans l'outil Word existant."""
        section = doc.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.header_distance = Cm(1.1)
        section.footer_distance = Cm(1.1)

        normal = doc.styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor(24, 39, 65)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for style_name, size in (("Title", 21), ("Subtitle", 12), ("Heading 1", 12)):
            style = doc.styles[style_name]
            style.font.name = "Aptos Display" if style_name != "Subtitle" else "Aptos"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(28, 72, 151) if style_name == "Heading 1" else RGBColor(24, 39, 65)
        heading_style = doc.styles["Heading 1"]
        heading_style.font.bold = True
        heading_style.paragraph_format.space_before = Pt(17)
        heading_style.paragraph_format.space_after = Pt(9)
        heading_style.paragraph_format.keep_with_next = True

        header = section.header.paragraphs[0]
        header.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
        header.add_run("RUINION AI").bold = True
        header.add_run("\tCOMPTE RENDU DE RÉUNION")
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(28, 72, 151)

        reference_match = re.match(r"meeting-(\d+)-", filename or "")
        reference = f"RAI-{int(reference_match.group(1)):05d}" if reference_match else "RUINION AI"
        footer = section.footer.paragraphs[0]
        footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
        footer.add_run(f"Réf. {reference}\tPage ")
        page_field = OxmlElement("w:fldSimple")
        page_field.set(qn("w:instr"), "PAGE")
        footer._p.append(page_field)
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(91, 105, 124)

        doc.core_properties.title = f"Compte rendu de réunion — {title}"
        doc.add_paragraph("COMPTE RENDU DE RÉUNION", style="Title")
        doc.add_paragraph(title or "Réunion", style="Subtitle")
        for label, value in (
            ("ORGANISATION", organization or "Non précisé"),
            ("DATE", date or "Non précisée"),
            ("PARTICIPANTS", participants or "Non précisé"),
            ("RÉFÉRENCE", reference),
        ):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            caption = paragraph.add_run(f"{label}   ")
            caption.bold = True
            caption.font.size = Pt(8)
            caption.font.color.rgb = RGBColor(28, 72, 151)
            paragraph.add_run(value)

        summary, sections = self._meeting_content(body or "")
        self._meeting_section(doc, "1. Synthèse")
        self._add_meeting_lines(doc, summary)
        self._meeting_section(doc, "2. Échanges et points clés")
        discussion = "\n".join(value for value in (sections.get("discussion"), sections.get("points")) if value)
        self._add_meeting_lines(doc, discussion)
        self._meeting_section(doc, "3. Décisions")
        self._add_meeting_lines(doc, sections.get("decisions", ""))
        self._meeting_section(doc, "4. Plan d’actions")
        actions = self._meeting_action_rows(sections.get("actions", ""))
        if actions:
            table = self._create_table(doc, {
                "headers": ["Action", "Responsable", "Échéance"],
                "rows": actions,
                "styles": {
                    "table_style": "Table Grid",
                    "header_style": {"bold": True, "background": "EAF1FC", "color": "1C4897"},
                    "column_widths": [3.5, 1.5, 1.5],
                },
            })
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        else:
            doc.add_paragraph("Non précisé dans la transcription.")

        for heading, key in (("5. Points en suspens", "questions"), ("6. Risques et blocages", "risks")):
            content = sections.get(key, "")
            if not self._meeting_missing(content):
                self._meeting_section(doc, heading)
                self._add_meeting_lines(doc, content)

    @staticmethod
    def _save_word(doc, filename):
        reports_dir = get_reports_dir()
        safe_filename = (
            "".join(c for c in (filename or "") if c.isalnum() or c in (" ", "-", "_")).rstrip()
            or "document"
        )
        file_path = reports_dir / f"{safe_filename}.docx"
        doc.save(file_path)
        abs_path = file_path.resolve()
        return f"Document Word créé avec succès : {abs_path}||{abs_path}"

    def forward(
        self,
        title: str,
        recipient: str = "",
        sender: str = "",
        date: str = "",
        subject: str = "",
        body: str = "",
        tables: str = "",
        filename: str = "document",
        style_config: str = "",
        template: str = "",
        organization: str = "",
        participants: str = "",
    ):
        try:
            doc = Document()

            if (template or "").lower() == "meeting_report" or (filename or "").startswith("meeting-"):
                self._render_meeting_report(doc, title, date, body, filename, organization, participants)
                return self._save_word(doc, filename)

            # === Configuration du style ===
            config = {}
            if style_config:
                try:
                    config = json.loads(style_config)
                except:
                    pass

            font_name = config.get("font", "Arial")
            font_size = config.get("font_size", 11)
            margins = config.get("margins", 2.5)
            line_spacing = config.get("line_spacing", 1.15)

            # Marges
            section = doc.sections[0]
            section.top_margin = Cm(margins)
            section.bottom_margin = Cm(margins)
            section.left_margin = Cm(margins)
            section.right_margin = Cm(margins)

            # Style global
            style = doc.styles["Normal"]
            style.font.name = font_name
            style.font.size = Pt(font_size)

            # Interligne
            paragraph_format = style.paragraph_format
            paragraph_format.line_spacing = line_spacing

            # === En-tête lettre (optionnel) ===
            if sender:
                p = doc.add_paragraph(sender)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            if recipient:
                doc.add_paragraph(recipient)

            if date:
                doc.add_paragraph(date)

            if subject:
                p = doc.add_paragraph()
                p.add_run("Objet : ").bold = True
                p.add_run(subject)
                doc.add_paragraph()

            # === Titre ===
            if title:
                title_p = doc.add_paragraph(title)
                title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_run = title_p.runs[0]
                title_run.font.size = Pt(16)
                title_run.bold = True
                doc.add_paragraph()

            # === Corps avec tableaux inline ===
            if body:
                self._parse_body_with_tables(doc, body)

            # === Tableaux externes (en fin de document) ===
            if tables:
                try:
                    tables_array = json.loads(tables)
                    for table_config in tables_array:
                        doc.add_paragraph()
                        if "title" in table_config:
                            p = doc.add_paragraph(table_config["title"])
                            p.runs[0].bold = True
                        self._create_table(doc, table_config)
                        doc.add_paragraph()
                except json.JSONDecodeError:
                    pass

            # === Formule de politesse (si format lettre) ===
            if recipient and sender:
                doc.add_paragraph()
                doc.add_paragraph(
                    "Je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées."
                )
                doc.add_paragraph()
                doc.add_paragraph(sender.split("\n")[0])  # Premier ligne = nom

            return self._save_word(doc, filename)

        except Exception as e:
            return f"✗ Erreur lors de la création du document : {str(e)}"


class BuildPDF(Tool):
    name = "BuildPDF"
    description = (
        "Génère un PDF professionnel avec titre, contenu, marges et styles optimisés."
    )

    inputs = {
        "name": {"type": "string", "description": "Nom du fichier PDF sans extension"},
        "title": {"type": "string", "description": "Titre du PDF"},
        "content": {"type": "string", "description": "Texte du PDF"},
    }
    output_type = "string"

    def forward(self, name: str, title: str, content: str) -> str:
        try:
            safe_name = (
                "".join(
                    char
                    for char in str(name)
                    if char.isalnum() or char in (" ", "-", "_")
                ).strip()
                or "document"
            )
            reports_dir = get_reports_dir()

            file_path = reports_dir / f"{safe_name}.pdf"
            styles = getSampleStyleSheet()
            style_title = styles["Title"]
            style_body = styles["BodyText"]

            doc = SimpleDocTemplate(
                str(file_path),
                pagesize=A4,
                leftMargin=20,
                rightMargin=20,
                topMargin=20,
                bottomMargin=20
            )

            story = [
                Paragraph(escape(title), style_title),
                Spacer(1, 12),
                Paragraph(escape(content).replace("\n", "<br/>"), style_body),
            ]

            doc.build(story)
            

            return f"PDF '{file_path.name}' genere avec succes.||{file_path.resolve()}"

        except Exception as e:
            return f"Erreur PDF : {str(e)}"
